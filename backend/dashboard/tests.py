import tempfile
from datetime import datetime, timedelta

from django.test import TestCase, override_settings
from django.utils import timezone

from .models import Campaign, QASyncWindow
from .qa_source import (
    _merge_intervals, _record_synced_window, _subtract_intervals, compute_coverage,
)


def d(day, hour=0, minute=0, second=0):
    return timezone.make_aware(datetime(2026, 9, day, hour, minute, second))


class DashboardTeamStatsTests(TestCase):
    """The Altitude BPO online dashboard's team stats API (dashboard/team_stats.py)."""

    def setUp(self):
        from . import team_stats
        from .models import DashboardTeam

        team_stats._cache.clear()
        DashboardTeam.objects.create(floor='Floor 1', display_name='Team A', source_team_names=['TeamA', 'TeamA2'], target=10, sort_order=1)
        DashboardTeam.objects.create(floor='Floor 1', display_name='Team B', source_team_names=['TeamB'], target=5, sort_order=2)
        DashboardTeam.objects.create(floor='Floor 1', display_name='Unmapped', source_team_names=[], target=7, sort_order=3)
        DashboardTeam.objects.create(floor='Floor 2', display_name='Team C', source_team_names=['TeamC'], target=20)
        DashboardTeam.objects.create(floor='Floor 2', display_name='Retired', source_team_names=['TeamX'], target=99, is_active=False)

        self.db = {
            # calls, sales, answered, talk_seconds, agents
            'TeamA': {'calls': 100, 'sales': 4, 'answered': 10, 'talk_seconds': 600.0, 'agents': 5},
            'TeamA2': {'calls': 50, 'sales': 1, 'answered': 10, 'talk_seconds': 1400.0, 'agents': 5},
            'TeamB': {'calls': 40, 'sales': 2, 'answered': 2, 'talk_seconds': 20.0, 'agents': 2},
            'TeamC': {'calls': 80, 'sales': 6, 'answered': 20, 'talk_seconds': 4000.0, 'agents': 8},
        }

        from rest_framework.test import APIClient
        from django.contrib.auth.models import User
        self.client = APIClient()
        self.client.force_authenticate(User.objects.create_user('dashboard-service'))

    def _get(self, floor):
        from unittest import mock
        with mock.patch('dashboard.team_stats.fetch_team_day_stats', return_value=self.db):
            return self.client.get('/api/dashboard/team-stats/', {'floor': floor})

    def test_floor_rows_and_weighted_total_average(self):
        data = self._get('Floor 1').json()
        rows = {t['team']: t for t in data['teams']}
        # Team A combines two source teams: 5 sales, (600+1400)/(5+5 live agents) = 200s
        self.assertEqual(rows['Team A']['current'], 5)
        self.assertEqual(rows['Team A']['shortfall'], 5)
        self.assertEqual(rows['Team A']['avg_talk_seconds'], 200.0)
        self.assertEqual(rows['Team A']['agents_live'], 10)
        self.assertEqual(rows['Team B']['avg_talk_seconds'], 10.0)
        # Unmapped rows show zeros / no average, and say they're unmapped.
        self.assertEqual(rows['Unmapped']['current'], 0)
        self.assertIsNone(rows['Unmapped']['avg_talk_seconds'])
        self.assertFalse(rows['Unmapped']['mapped'])
        # Floor total: weighted, NOT the mean of 200 and 10 (= 105): 2020s / 12 agents = 168.3
        t = data['totals']
        self.assertEqual((t['target'], t['current'], t['shortfall']), (22, 7, 15))
        self.assertEqual(t['avg_talk_seconds'], 168.3)
        self.assertEqual(t['agents_live'], 12)
        self.assertEqual(t['calls_answered'], 22)

    def test_team_pat_pinned_first_then_by_sales(self):
        from .models import DashboardTeam
        DashboardTeam.objects.create(floor='Floor 2', display_name='Team Pat', source_team_names=['TeamP'], target=5, sort_order=9)
        self.db['TeamP'] = {'calls': 1, 'sales': 0, 'answered': 0, 'talk_seconds': 0.0, 'agents': 1}
        self.assertEqual([t['team'] for t in self._get('Floor 2').json()['teams']], ['Team Pat', 'Team C'])
        # Floor 1: no Pat, so purely most sales first (Team A 5, Team B 2, Unmapped 0).
        self.assertEqual([t['team'] for t in self._get('Floor 1').json()['teams']], ['Team A', 'Team B', 'Unmapped'])

    def test_floor_is_isolated_and_inactive_rows_hidden(self):
        data = self._get('Floor 2').json()
        self.assertEqual([t['team'] for t in data['teams']], ['Team C'])
        self.assertEqual(data['totals']['target'], 20)

    def test_global_combines_both_floors(self):
        data = self._get('Global').json()
        self.assertEqual(len(data['teams']), 4)
        self.assertEqual(data['totals']['current'], 13)
        self.assertEqual(data['totals']['avg_talk_seconds'], round((2020 + 4000) / (12 + 8), 1))

    def test_unknown_floor_and_db_failure_and_auth(self):
        from unittest import mock
        from rest_framework.test import APIClient
        from .external_source import ExternalSourceError

        self.assertEqual(self._get('Floor 9').status_code, 400)
        with mock.patch('dashboard.team_stats.fetch_team_day_stats', side_effect=ExternalSourceError('db down')):
            self.assertEqual(self.client.get('/api/dashboard/team-stats/', {'floor': 'Floor 1'}).status_code, 502)
        self.assertEqual(APIClient().get('/api/dashboard/team-stats/').status_code, 401)

    def test_target_can_be_edited(self):
        resp = self.client.post('/api/dashboard/team-targets/', {'floor': 'Floor 1', 'team': 'Team A', 'target': 25}, format='json')
        self.assertEqual(resp.status_code, 200)
        from .models import DashboardTeam
        self.assertEqual(DashboardTeam.objects.get(display_name='Team A').target, 25)
        for bad in ({'target': -1}, {'target': 'abc'}):
            r = self.client.post('/api/dashboard/team-targets/', {'floor': 'Floor 1', 'team': 'Team A', **bad}, format='json')
            self.assertEqual(r.status_code, 400)
        r = self.client.post('/api/dashboard/team-targets/', {'floor': 'Floor 1', 'team': 'Nope', 'target': 1}, format='json')
        self.assertEqual(r.status_code, 404)


class DashboardTokenCommandTests(TestCase):
    def test_token_comes_from_the_environment_variable_and_works_for_the_api(self):
        from unittest import mock

        from django.core.management import call_command
        from rest_framework.test import APIClient

        token = 'a1b2c3d4e5f60718293a4b5c6d7e8f9012345678'
        with mock.patch.dict('os.environ', {'DASHBOARD_API_TOKEN': token}):
            call_command('dashboard_api_token', '--from-env-only')
            call_command('dashboard_api_token', '--from-env-only')   # re-running is a no-op

            client = APIClient()
            client.credentials(HTTP_AUTHORIZATION=f'Token {token}')
            self.assertEqual(client.get('/api/dashboard/team-stats/', {'floor': 'Floor 9'}).status_code, 400)  # authenticated

            # Changing the variable swaps the token: the old one stops working.
            new = 'ffffffffffffffffffffffffffffffffffffffff'
            with mock.patch.dict('os.environ', {'DASHBOARD_API_TOKEN': new}):
                call_command('dashboard_api_token', '--from-env-only')
            self.assertEqual(client.get('/api/dashboard/team-stats/').status_code, 401)

    def test_from_env_only_does_nothing_when_unset_and_rejects_bad_lengths(self):
        from unittest import mock

        from django.core.management import call_command
        from django.core.management.base import CommandError
        from rest_framework.authtoken.models import Token

        with mock.patch.dict('os.environ', {'DASHBOARD_API_TOKEN': ''}):
            call_command('dashboard_api_token', '--from-env-only')
            self.assertEqual(Token.objects.count(), 0)
        with mock.patch.dict('os.environ', {'DASHBOARD_API_TOKEN': 'x' * 41}):
            with self.assertRaises(CommandError):
                call_command('dashboard_api_token', '--from-env-only')


class SeedDashboardTeamsTests(TestCase):
    def test_floor_2_teams_are_seeded_mapped_and_reruns_keep_targets(self):
        from django.core.management import call_command

        from .models import DashboardTeam

        # A leftover unmapped placeholder from the earlier seed is cleaned up.
        DashboardTeam.objects.create(floor='Floor 2', display_name='Team Victus Funeral', source_team_names=[], target=30)

        call_command('seed_dashboard_teams')
        rows = {r.display_name: r for r in DashboardTeam.objects.filter(floor='Floor 2')}
        self.assertEqual(list(rows), ['Team Themba', 'Team Thelma', 'Team Ayanda', 'Team Sandra', 'Team Pat'])
        self.assertEqual(rows['Team Themba'].source_team_names, ['TeamThemba'])
        self.assertEqual(rows['Team Sandra'].source_team_names, ['Team_Sandra', 'Team Sandra'])
        self.assertIn('VodacomFuneralUpsell', rows['Team Pat'].source_team_names)
        self.assertTrue(all(r.target == 0 for r in rows.values()))

        rows['Team Pat'].target = 60
        rows['Team Pat'].save()
        call_command('seed_dashboard_teams')   # re-running must not reset a target someone set
        self.assertEqual(DashboardTeam.objects.get(display_name='Team Pat').target, 60)
        self.assertEqual(DashboardTeam.objects.filter(floor='Floor 2').count(), 5)

    def test_floor_1_teams_1life_is_one_team_from_both_source_teams(self):
        from django.core.management import call_command

        from .models import DashboardTeam

        call_command('seed_dashboard_teams')
        rows = {r.display_name: r for r in DashboardTeam.objects.filter(floor='Floor 1')}
        self.assertEqual(list(rows), ['1Life', 'Assupol Tactical', 'Assupol Mighty', 'Hollard'])
        self.assertEqual(rows['1Life'].source_team_names, ['1Life', '1 Life'])           # one row, two source teams
        self.assertEqual(rows['Assupol Tactical'].source_team_names, ['Assupol Tectical Unit'])
        self.assertEqual(rows['Assupol Mighty'].source_team_names, ['Mighty Assupol'])
        self.assertEqual(rows['Hollard'].source_team_names, ['Hollard Edgars'])
        self.assertNotIn('Assupol', sum((r.source_team_names for r in rows.values()), []))  # the plain 1-agent team is excluded


class IntervalMathTests(TestCase):
    def test_merge_joins_overlapping_and_touching(self):
        merged = _merge_intervals([
            (d(17), d(18, 23, 59, 59)),
            (d(19), d(20)),            # starts 1s after the previous end: touching
            (d(25), d(26)),
        ])
        self.assertEqual(merged, [(d(17), d(20)), (d(25), d(26))])

    def test_subtract_finds_every_gap(self):
        gaps = _subtract_intervals((d(17), d(30)), [(d(20), d(22)), (d(25), d(27))])
        self.assertEqual(gaps, [(d(17), d(20)), (d(22), d(25)), (d(27), d(30))])

    def test_subtract_fully_covered_is_empty(self):
        self.assertEqual(_subtract_intervals((d(20), d(22)), [(d(17), d(30))]), [])

    def test_subtract_no_cuts_is_whole_range(self):
        self.assertEqual(_subtract_intervals((d(20), d(22)), []), [(d(20), d(22))])


class CoverageTests(TestCase):
    def setUp(self):
        self.campaign = Campaign.objects.create(name='Absa Insurance', display_name='Absa Insurance')

    def test_unsynced_campaign_is_one_big_gap(self):
        gaps = compute_coverage(self.campaign, d(17), d(30, 23, 59, 59))
        self.assertEqual(gaps, [(d(17), d(30, 23, 59, 59))])

    def test_partial_sync_reports_the_missing_front(self):
        # The real incident: only 29-30 Sep synced, view asks for 17-30 Sep.
        _record_synced_window(self.campaign, d(29), d(30, 23, 59, 59))
        gaps = compute_coverage(self.campaign, d(17), d(30, 23, 59, 59))
        self.assertEqual(gaps, [(d(17), d(29))])

    def test_syncing_the_gap_closes_it_and_merges_windows(self):
        _record_synced_window(self.campaign, d(29), d(30, 23, 59, 59))
        _record_synced_window(self.campaign, d(17), d(28, 23, 59, 59))
        self.assertEqual(compute_coverage(self.campaign, d(17), d(30, 23, 59, 59)), [])
        self.assertEqual(QASyncWindow.objects.filter(campaign=self.campaign).count(), 1)

    def test_skipped_ranges_stay_flagged(self):
        _record_synced_window(self.campaign, d(17), d(30, 23, 59, 59), skipped=[(d(22), d(23))])
        gaps = compute_coverage(self.campaign, d(17), d(30, 23, 59, 59))
        self.assertEqual(gaps, [(d(22), d(23))])

    def test_records_endpoint_returns_gaps_ready_to_post_back_to_sync(self):
        from django.contrib.auth.models import User
        from rest_framework.test import APIClient

        _record_synced_window(self.campaign, d(29), d(30, 23, 59, 59))
        client = APIClient()
        client.force_authenticate(User.objects.create_user('qa-tester'))
        resp = client.get('/api/qa/records/', {
            'campaign_ids': str(self.campaign.id), 'start_date': '2026-09-17', 'end_date': '2026-09-30',
        })
        self.assertEqual(resp.status_code, 200)
        coverage = resp.json()['coverage']
        self.assertEqual(len(coverage), 1)
        self.assertEqual(coverage[0]['campaign'], 'Absa Insurance')
        gap = coverage[0]['gaps'][0]
        self.assertEqual((gap['start_date'], gap['start_time']), ('2026-09-17', '00:00:00'))
        self.assertEqual(gap['end_date'], '2026-09-29')

    def test_records_endpoint_has_no_coverage_for_open_ended_filter(self):
        from django.contrib.auth.models import User
        from rest_framework.test import APIClient

        client = APIClient()
        client.force_authenticate(User.objects.create_user('qa-tester2'))
        resp = client.get('/api/qa/records/', {'campaign_ids': str(self.campaign.id)})
        self.assertEqual(resp.json()['coverage'], [])

    def test_download_has_lead_id_and_the_extra_columns(self):
        from io import BytesIO

        import openpyxl
        from django.contrib.auth.models import User
        from rest_framework.test import APIClient

        from .models import QACallRecord

        QACallRecord.objects.create(
            campaign=self.campaign, interaction_id='i-1', contact_id='16601553', customer='MOKEBE MOGOWE',
            phone_number='+27823214056', alt_phone_number='+27781940673', id_number='8610160754089',
            batch_name='Absa QA', agent_name='Agent A', outcome='Sale Made - Vehicle Insurance',
            call_date=d(30, 10, 5), call_end=d(30, 10, 9), direction='outbound', talk_seconds=240,
            recording_key='rec_1', recording_duration_seconds=238, extension='ComLen1', hangup_user='agent',
        )
        client = APIClient()
        user = User.objects.create_user('qa-dl')
        client.force_authenticate(user)
        with tempfile.TemporaryDirectory() as media, override_settings(MEDIA_ROOT=media):
            resp = client.get('/api/qa/download/', {'campaign_ids': str(self.campaign.id)})
            self.assertEqual(resp.status_code, 200)
            ws = openpyxl.load_workbook(BytesIO(resp.content)).active
            header = [c.value for c in ws[1]]
            row = dict(zip(header, [c.value for c in ws[2]]))
            for needed in ('Lead ID', 'Customer', 'Data List', 'Phone Number', 'Agent Name', 'Extension',
                           'Outcome', 'Campaign', 'Hangup User', 'Lead Reference', 'ID Number',
                           'Alt Phone Number', 'Interaction ID'):
                self.assertIn(needed, header)
            self.assertEqual(row['Lead ID'], '16601553')
            self.assertEqual(row['ID Number'], '8610160754089')   # text, not a number
            self.assertEqual(row['Data List'], 'Absa QA')
            self.assertEqual(row['Extension'], 'ComLen1')
            self.assertEqual(row['Hangup User'], 'agent')
            self.assertEqual(row['Talk Time (seconds)'], 240)
            self.assertEqual(row['Call Start (UTC)'], '2026-09-30 10:05:00')
            self.assertIsNone(row['Lead Reference'])               # blank stays blank

            # The download shows up in Recent activity, and its file can be fetched again.
            activity = client.get('/api/qa/activity/').json()
            self.assertEqual(len(activity), 1)
            self.assertEqual(activity[0]['kind'], 'download')
            self.assertEqual(activity[0]['status'], 'done')
            self.assertEqual(activity[0]['records'], 1)
            self.assertEqual(activity[0]['campaign'], 'Absa Insurance')
            self.assertEqual(activity[0]['filters']['campaign_ids'], [self.campaign.id])
            self.assertTrue(activity[0]['has_file'])
            again = client.get(f"/api/qa/activity/{activity[0]['id']}/file/")
            self.assertEqual(again.status_code, 200)
            self.assertEqual(b''.join(again.streaming_content), resp.content)
            again.close()  # release the file handle (matters for cleaning up the temp dir on Windows)

            # Someone else can't see or fetch it.
            other = APIClient()
            other.force_authenticate(User.objects.create_user('someone-else'))
            self.assertEqual(other.get('/api/qa/activity/').json(), [])
            self.assertEqual(other.get(f"/api/qa/activity/{activity[0]['id']}/file/").status_code, 404)

    def test_sync_endpoint_logs_recent_activity_for_success_and_failure(self):
        from unittest import mock

        from django.contrib.auth.models import User
        from rest_framework.test import APIClient

        from .external_source import ExternalSourceError

        client = APIClient()
        client.force_authenticate(User.objects.create_user('qa-sync'))
        body = {'campaign_ids': [self.campaign.id], 'start_date': '2026-09-17', 'end_date': '2026-09-30'}

        with mock.patch('dashboard.qa_source.sync_campaign_qa_cache', return_value=(412841, timezone.now())):
            self.assertEqual(client.post('/api/qa/sync/', body, format='json').status_code, 200)
        with mock.patch('dashboard.qa_source.sync_campaign_qa_cache', side_effect=ExternalSourceError('db down')):
            client.post('/api/qa/sync/', body, format='json')

        activity = client.get('/api/qa/activity/').json()
        self.assertEqual([(a['kind'], a['status']) for a in activity], [('sync', 'failed'), ('sync', 'done')])
        self.assertEqual(activity[1]['records'], 412841)
        self.assertEqual(activity[0]['message'], 'db down')
        self.assertEqual(activity[1]['filters']['start_date'], '2026-09-17')

    def test_stop_cancels_a_running_sync_and_records_no_window(self):
        from unittest import mock

        from . import qa_source
        from .external_source import ExternalSourceError

        self.campaign.cd_campaign_id = 'b45ceab8-2ce0-4114-a724-5f11cc3d3231'
        self.campaign.save()
        row = {
            'interaction_id': 'i-1', 'contact_id': 1, 'customer_id': 'c', 'firstname': 'A', 'lastname': 'B',
            'phone_number': '1', 'alt_phone_number': None, 'lead_reference': None, 'id_number': None,
            'agent_name': 'x', 'outcome': 'Sale', 'call_date': d(30), 'call_end': None, 'direction': None,
            'talk_time': None, 'recording_key': None, 'recording_duration': None, 'batch_name': None,
        }

        def fake_fetch(cd_id, start, end, on_window=None, skipped_out=None):
            # The user hits Stop while the first window is being fetched; the
            # real fetch wraps callback errors in ExternalSourceError.
            self.assertTrue(qa_source.request_cancel(self.campaign.id))
            try:
                on_window([row])
            except Exception as e:
                raise ExternalSourceError(f'wrapped: {e}')

        with mock.patch.object(qa_source, 'fetch_qa_interactions', fake_fetch):
            with self.assertRaises(qa_source.SyncCancelled):
                qa_source.sync_campaign_qa_cache(self.campaign, start_date='2026-09-29', end_date='2026-09-30')

        self.assertEqual(QASyncWindow.objects.filter(campaign=self.campaign).count(), 0)
        self.assertFalse(qa_source.request_cancel(self.campaign.id))  # nothing running, flags cleared

    def test_future_part_of_range_is_not_a_gap(self):
        now = timezone.now()
        _record_synced_window(self.campaign, now - timedelta(days=2), now)
        gaps = compute_coverage(self.campaign, now - timedelta(days=2), now + timedelta(days=5))
        self.assertEqual(gaps, [])
